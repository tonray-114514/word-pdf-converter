"""体检：确认自带 env 的关键能力与转换功能是否正常。

不依赖网络，全部离线自检。
"""
import os
import sys
import traceback

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, APP)

ok = []


def check(name, fn):
    try:
        result = fn()
        print(f"通过  {name}" + (f"  [{result}]" if result else ""))
        ok.append(True)
    except Exception as exc:  # noqa: BLE001
        print(f"失败  {name}  -> {type(exc).__name__}: {exc}")
        traceback.print_exc()
        ok.append(False)


print("Python:", sys.version.split()[0], "|", sys.executable)
print("=" * 64)


def _ssl():
    import ssl

    return ssl.OPENSSL_VERSION


def _hashing():
    import hashlib

    return hashlib.sha256(b"x").hexdigest()[:12]


def _zlib():
    import zlib

    return f"crc32={zlib.crc32(b'x')}"


def _sqlite():
    import sqlite3

    return sqlite3.sqlite_version


def _tk():
    import tkinter

    return f"Tk {tkinter.TkVersion}"


def _docx():
    import docx

    d = docx.Document()
    d.add_paragraph("测试")
    return "可读写"


def _pymupdf():
    import pymupdf

    return pymupdf.__doc__.split(":")[0]


def _pywin32():
    import pythoncom
    import win32com.client  # noqa: F401

    return os.path.basename(pythoncom.__file__)


def _converter():
    import converter

    return converter.describe_engine()


def _mkdtemp():
    """临时目录是否可用（PDF→Word 生成图片时要用）。"""
    import uuid

    import converter

    root = converter.temp_root()
    path = os.path.join(root, "check_" + uuid.uuid4().hex[:8])
    os.makedirs(path, exist_ok=True)
    try:
        with open(os.path.join(path, "a.txt"), "w", encoding="utf-8") as fh:
            fh.write("ok")
        with open(os.path.join(path, "a.txt"), encoding="utf-8") as fh:
            if fh.read() != "ok":
                raise OSError("写入内容不一致")
        return root
    finally:
        import shutil

        shutil.rmtree(path, ignore_errors=True)


def _pdf2word():
    """完整跑一遍 PDF→Word（不需要 Word，纯 Python 路径）。"""
    import pymupdf
    import converter

    src = os.path.join(APP, "_check_in.pdf")
    dst = os.path.join(APP, "_check_out.docx")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Check Heading", fontsize=18, fontname="hebo")
    page.insert_text((72, 130), "Body text 12345.", fontsize=11)
    doc.save(src)
    doc.close()
    try:
        res = converter.pdf_to_word(src, dst)
        import docx

        d = docx.Document(dst)
        return f"{res.pages} 页 -> {len(d.paragraphs)} 段落"
    finally:
        for p in (src, dst):
            if os.path.exists(p):
                os.remove(p)


def _batfiles():
    """检查 .bat 的编码与换行符（错了会导致双击无法启动）。"""
    names = sorted(n for n in os.listdir(APP) if n.lower().endswith(".bat"))
    if not names:
        raise OSError("未找到任何 .bat 文件")
    bad = []
    for name in names:
        with open(os.path.join(APP, name), "rb") as fh:
            data = fh.read()
        if data[:3] == b"\xef\xbb\xbf":
            bad.append(f"{name}(有BOM)")
            continue
        if data.count(b"\r") != data.count(b"\n"):
            bad.append(f"{name}(换行非CRLF)")
            continue
        try:
            data.decode("gbk")
        except UnicodeDecodeError:
            bad.append(f"{name}(非GBK编码)")
    if bad:
        raise OSError("格式有问题的批处理: " + ", ".join(bad))
    return f"{len(names)} 个批处理格式正确"


check("ssl（HTTPS 基础）", _ssl)
check("hashlib", _hashing)
check("zlib", _zlib)
check("sqlite3", _sqlite)
check("tkinter（桌面界面）", _tk)
check("python-docx", _docx)
check("PyMuPDF", _pymupdf)
check("pywin32（Word 引擎）", _pywin32)
check("转换引擎探测", _converter)
check("临时目录可写", _mkdtemp)
check("启动脚本格式（CRLF+GBK）", _batfiles)
check("PDF → Word 全流程", _pdf2word)

print("=" * 64)
print(f"通过 {sum(ok)}/{len(ok)}")
sys.exit(0 if all(ok) else 1)
