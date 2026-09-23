"""手机网页版：端到端 + 历史记录 + 引擎信息 测试。

用法: python test_web2.py [服务地址]
"""
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

APP = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SCRATCH = os.path.join(APP, ".tmp")
os.makedirs(SCRATCH, exist_ok=True)
BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8793"
DOCX = os.path.join(SCRATCH, "空格与中文 文件名.docx")
PDF = os.path.join(SCRATCH, "下载结果.pdf")

results = []


def check(name, cond, extra=""):
    print(("通过  " if cond else "失败  ") + name + (f"  {extra}" if extra else ""))
    results.append(bool(cond))


def get_json(path):
    with urllib.request.urlopen(BASE + path, timeout=15) as r:
        return r.status, json.load(r)


def post_file(path, direction, token):
    data = open(path, "rb").read()
    qs = urllib.parse.urlencode({"token": token})
    req = urllib.request.Request(
        f"{BASE}/api/convert?{qs}", data=data, method="POST",
        headers={"X-Filename": urllib.parse.quote(os.path.basename(path)),
                 "X-Direction": direction,
                 "Content-Length": str(len(data))})
    with urllib.request.urlopen(req, timeout=180) as r:
        return json.load(r)


def wait(job_id, timeout=180):
    t0 = time.time()
    while time.time() - t0 < timeout:
        _, d = get_json(f"/api/status?id={job_id}")
        if d["status"] in ("完成", "失败", "已取消"):
            return d
        time.sleep(0.7)
    raise TimeoutError("超时")


def make_docx():
    import docx

    d = docx.Document()
    d.add_heading("手机端测试文档", level=0)
    d.add_paragraph("用于验证手机网页版的完整链路。")
    d.add_paragraph("含中文与 English 1234567890。")
    d.save(DOCX)


print("=== 1. 服务与引擎信息 ===")
code, info = get_json("/api/info")
check("/api/info 可访问", code == 200)
check("返回引擎描述", bool(info.get("engine")), info.get("engine", "")[:40])
check("返回 engine_mode", "engine_mode" in info, str(info.get("engine_mode")))
check("返回 can_word2pdf", "can_word2pdf" in info, str(info.get("can_word2pdf")))
check("本机 Word→PDF 可用", info.get("can_word2pdf") is True)

print("\n=== 2. 生成测试文件（含空格与中文文件名）===")
make_docx()
print(f"   {os.path.basename(DOCX)} ({os.path.getsize(DOCX)} bytes)")

print("\n=== 3. Word → PDF ===")
job = post_file(DOCX, "word2pdf", "tk_alpha")
check("任务已创建", job.get("status") == "排队中", job.get("status", ""))
res = wait(job["id"])
check("转换完成", res["status"] == "完成", res.get("error", ""))
check("返回下载链接", bool(res.get("download")))
check("结果文件名正确",
      res.get("result_name", "").startswith("空格与中文"), res.get("result_name", ""))

with urllib.request.urlopen(BASE + res["download"], timeout=60) as r:
    body = r.read()
with open(PDF, "wb") as fh:
    fh.write(body)
check("下载内容是 PDF", body[:5] == b"%PDF-", f"{len(body)} bytes")

print("\n=== 4. PDF → Word（引擎参数链路）===")
job2 = post_file(PDF, "pdf2word", "tk_beta")
res2 = wait(job2["id"])
check("反向转换完成", res2["status"] == "完成", res2.get("error", ""))
if res2["status"] == "完成":
    with urllib.request.urlopen(BASE + res2["download"], timeout=60) as r:
        b2 = r.read()
    check("下载内容是 docx", b2[:2] == b"PK", f"{len(b2)} bytes")

print("\n=== 5. 最近转换记录 ===")
code, rec = get_json("/api/recent")
check("/api/recent 可访问", code == 200)
jobs = rec.get("jobs", [])
check("记录非空", len(jobs) >= 2, f"{len(jobs)} 条")
done = [j for j in jobs if j["status"] == "完成" and j.get("download")]
check("完成记录带下载链接", len(done) >= 2, f"{len(done)} 条可下载")
names = [j.get("result_name", "") for j in done]
check("记录含本次结果", any("空格与中文" in n for n in names), str(names[:3]))

print("\n=== 6. 口令外仍可读到历史（本次服务未设口令）===")
check("未设口令时不要求登录", info.get("need_password") is False)

print()
print(f"通过 {sum(results)}/{len(results)}")
sys.exit(0 if all(results) else 1)
